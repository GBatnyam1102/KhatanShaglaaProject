import argparse
from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path("docs/Khatan_Shaglaa_Website_Requirements.md")
OUTPUT = Path("docs/Khatan_Shaglaa_Website_Requirements.docx")
DEFAULT_TITLE = "Хатан шаглаа web site"
SUBTITLE = "Функциональ шаардлага, функциональ бус шаардлага, архитектурын санал"
FONT_NAME = "Arial"
BRAND = RGBColor(111, 78, 55)
GOLD = RGBColor(184, 134, 11)
DARK = RGBColor(30, 24, 20)


def set_rfonts(element, name=FONT_NAME):
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def strip_inline(text):
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    return text.strip()


def add_inline(paragraph, text, size=None):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        value = value.replace("`", "")
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.name = FONT_NAME
        if size:
            run.font.size = Pt(size)
        set_rfonts(run._element)


def split_table_row(line):
    return [strip_inline(cell) for cell in line.strip().strip("|").split("|")]


def add_table(document, table_lines):
    header = split_table_row(table_lines[0])
    rows = [split_table_row(line) for line in table_lines[2:]]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.autofit = True
    for index, value in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, "E8D8B0")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = FONT_NAME
                run.font.size = Pt(8.5)
                set_rfonts(run._element)
    for row in rows:
        cells = table.add_row().cells
        for index in range(len(header)):
            value = row[index] if index < len(row) else ""
            cells[index].text = value
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_NAME
                    run.font.size = Pt(8)
                    set_rfonts(run._element)
    document.add_paragraph()


def add_code_block(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    for line in lines:
        run = paragraph.add_run(line + "\n")
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        set_rfonts(run._element, "Courier New")


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        set_rfonts(style._element)
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Title"].font.size = Pt(22)
    document.styles["Title"].font.color.rgb = DARK
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 1"].font.color.rgb = BRAND
    document.styles["Heading 2"].font.size = Pt(12.5)
    document.styles["Heading 2"].font.color.rgb = BRAND
    document.styles["Heading 3"].font.size = Pt(11)
    document.styles["Heading 3"].font.color.rgb = BRAND


def build_docx():
    document = Document()
    configure_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    first_heading = next((line[2:].strip() for line in lines if line.startswith("# ")), DEFAULT_TITLE)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(first_heading)
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(22)
    run.font.color.rgb = DARK
    set_rfonts(run._element)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run(SUBTITLE)
    srun.font.name = FONT_NAME
    srun.font.size = Pt(12)
    srun.font.color.rgb = GOLD
    set_rfonts(srun._element)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(meta, f"Огноо: {date.today().isoformat()}")
    document.add_page_break()

    index = 0
    in_code = False
    code_lines = []
    while index < len(lines):
        line = lines[index].rstrip()
        if index == 0 and line.startswith("# "):
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|\s*[-:]+", lines[index + 1]):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            document.add_heading(strip_inline(heading_match.group(2)), level=level)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            index += 1
            continue
        numbered_match = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered_match.group(1))
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, line)
        index += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT.as_posix()}")


def parse_args():
    parser = argparse.ArgumentParser(description="Convert markdown to docx with brand styling.")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--subtitle", default=SUBTITLE)
    parser.add_argument("--default-title", dest="default_title", default=DEFAULT_TITLE)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    SOURCE = args.source
    OUTPUT = args.output
    SUBTITLE = args.subtitle
    DEFAULT_TITLE = args.default_title
    build_docx()
